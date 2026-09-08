"""Write a sweep result set as an .xlsx workbook with charts.

Layout mirrors the reference file `Aquisições_Dimming.xlsx`:

  * a **Dados** sheet -- the full flat table (every `harness.Row` column),
    one row per measured point;
  * one sheet per linearization state (**NãoLinearizado** / **Linearizado**)
    holding the `Modo dim | Dimming (%) | Fonte 1 (W) | Fonte 2 (W) | Total`
    table -- `Total` is a live `=C+D` formula -- and a smooth-line scatter
    chart of Total vs Dimming %, one series per driver mode.

Only the linearization sheets that actually have data are created, so a
plain single sweep produces Dados + one summary sheet.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable, Sequence

from openpyxl import Workbook
from openpyxl.chart import Reference, ScatterChart, Series
from openpyxl.chart.marker import Marker
from openpyxl.worksheet.worksheet import Worksheet

SUMMARY_HEADERS = ["Modo dim", "Dimming (%)", "Fonte 1 (W)", "Fonte 2 (W)", "Total"]

_LIN_SHEET = {False: "NãoLinearizado", True: "Linearizado"}
_MODE_ORDER = ["hybrid", "analog", "pwm"]
_MODE_LABEL = {"hybrid": "Híbrido", "analog": "Analógico", "pwm": "PWM"}


def _mode_key(name: str) -> int:
    return _MODE_ORDER.index(name) if name in _MODE_ORDER else len(_MODE_ORDER)


def write_report(
    rows: Sequence[dict[str, Any]],
    path: str | Path,
    *,
    flat_fields: Iterable[str],
) -> Path:
    """Build the workbook for `rows` and save it to `path`; returns Path."""
    wb = Workbook()
    # `Total` is a formula column -- make Excel recalc it (and the charts
    # that read it) on open.
    wb.calculation.fullCalcOnLoad = True
    _dados_sheet(wb.active, rows, list(flat_fields))

    by_lin: dict[bool, list[dict[str, Any]]] = {}
    for r in rows:
        by_lin.setdefault(bool(r["lin_enable"]), []).append(r)

    for lin in (False, True):
        if by_lin.get(lin):
            _summary_sheet(wb.create_sheet(_LIN_SHEET[lin]), by_lin[lin])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _dados_sheet(ws: Worksheet, rows: Sequence[dict[str, Any]], fields: list[str]) -> None:
    ws.title = "Dados"
    ws.append(fields)
    for r in rows:
        ws.append([r.get(f) for f in fields])
    ws.freeze_panes = "A2"


def _summary_sheet(ws: Worksheet, rows: Sequence[dict[str, Any]]) -> None:
    ws.append(SUMMARY_HEADERS)

    modes = sorted({r["mode_name"] for r in rows}, key=_mode_key)
    blocks: list[tuple[str, int, int]] = []
    xl = 1  # last written Excel row (row 1 = header)

    for mode in modes:
        label = _MODE_LABEL.get(mode, mode)
        start = xl + 1
        for r in [x for x in rows if x["mode_name"] == mode]:
            xl += 1
            ws.cell(xl, 1, label)
            ws.cell(xl, 2, r["dimming_pct"]).number_format = "0"
            ws.cell(xl, 3, round(float(r["f1_w"]), 3)).number_format = "0.00"
            ws.cell(xl, 4, round(float(r["f2_w"]), 3)).number_format = "0.00"
            ws.cell(xl, 5, f"=C{xl}+D{xl}").number_format = "0.00"
        if xl >= start:
            blocks.append((label, start, xl))

    for col, width in (("A", 14), ("B", 12), ("C", 12), ("D", 12), ("E", 10)):
        ws.column_dimensions[col].width = width

    _add_chart(ws, blocks)


def _add_chart(ws: Worksheet, blocks: list[tuple[str, int, int]]) -> None:
    if not blocks:
        return
    ch = ScatterChart()
    ch.scatterStyle = "smoothMarker"
    ch.x_axis.title = "Dimming (%)"
    ch.y_axis.title = "Potência (W)"
    ch.x_axis.delete = False
    ch.y_axis.delete = False
    ch.x_axis.axPos = "b"
    ch.y_axis.axPos = "l"
    ch.height = 11
    ch.width = 18

    for label, r0, r1 in blocks:
        xref = Reference(ws, min_col=2, min_row=r0, max_row=r1)
        yref = Reference(ws, min_col=5, min_row=r0, max_row=r1)
        s = Series(yref, xref, title=label)
        s.smooth = True
        s.marker = Marker(symbol="circle", size=5)
        ch.series.append(s)

    ws.add_chart(ch, "G2")
